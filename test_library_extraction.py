"""
Quick local test for the new Library file-type extraction (.html, .md, .pptx).
No Azure services needed — builds small in-memory samples and calls
sharepoint_client._extract_text() directly.

Usage:
    python test_library_extraction.py
"""

import io

from sharepoint_client import _extract_text

FAILURES = []


def check(label: str, text: str, must_contain: list[str], must_not_contain: list[str]):
    ok = True
    if not text.strip():
        print(f"[FAIL] {label}: extracted text is empty")
        ok = False
    for phrase in must_contain:
        if phrase not in text:
            print(f"[FAIL] {label}: expected to find {phrase!r} in extracted text")
            ok = False
    for phrase in must_not_contain:
        if phrase in text:
            print(f"[FAIL] {label}: raw syntax {phrase!r} leaked into extracted text")
            ok = False
    if ok:
        print(f"[PASS] {label}")
    else:
        FAILURES.append(label)
    print(f"  --- extracted text ---\n  {text!r}\n")


# --------------------------------------------------------------------------
# .html — trafilatura should keep the article, drop nav/footer boilerplate
# --------------------------------------------------------------------------

html_sample = """
<html>
<head><title>Test Page</title></head>
<body>
  <nav>Home | About | Contact</nav>
  <header>Site Header Banner</header>
  <article>
    <h1>VPN Troubleshooting Guide</h1>
    <p>If your VPN connection fails, first check that your internet connection
    is stable. Then verify your VPN client is up to date and restart it.</p>
    <p>If the issue persists, contact IT support with your error code.</p>
  </article>
  <footer>Copyright 2026 - All rights reserved - Privacy Policy</footer>
</body>
</html>
"""
html_text = _extract_text("guide.html", html_sample.encode("utf-8"))
check(
    "html",
    html_text,
    must_contain=["VPN", "internet connection is stable"],
    must_not_contain=["<article>", "<p>", "Home | About | Contact"],
)

# --------------------------------------------------------------------------
# .md — markdown syntax should be stripped, not just passed through raw
# --------------------------------------------------------------------------

md_sample = """# Password Reset Policy

Passwords **must** be changed every 90 days.

- Minimum length: 12 characters
- Must include a number and a symbol

See the [IT portal](https://example.com/it) for self-service reset.
"""
md_text = _extract_text("policy.md", md_sample.encode("utf-8"))
check(
    "md",
    md_text,
    must_contain=["Password Reset Policy", "must", "Minimum length: 12 characters"],
    must_not_contain=["**", "##", "- Minimum", "[IT portal]"],
)

# --------------------------------------------------------------------------
# .pptx — built in-memory with python-pptx itself, then extracted back out
# --------------------------------------------------------------------------

from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
slide_layout = prs.slide_layouts[1]  # title + content

slide1 = prs.slides.add_slide(slide_layout)
slide1.shapes.title.text = "Onboarding Overview"
slide1.placeholders[1].text_frame.text = "Welcome to the IT onboarding process."

slide2 = prs.slides.add_slide(slide_layout)
slide2.shapes.title.text = "Laptop Setup"
body = slide2.placeholders[1].text_frame
body.text = "Connect to the corporate VPN before installing software."
p2 = body.add_paragraph()
p2.text = "Contact the helpdesk if setup fails."

buf = io.BytesIO()
prs.save(buf)
pptx_bytes = buf.getvalue()

pptx_text = _extract_text("onboarding.pptx", pptx_bytes)
check(
    "pptx",
    pptx_text,
    must_contain=[
        "Onboarding Overview",
        "Welcome to the IT onboarding process.",
        "Laptop Setup",
        "Connect to the corporate VPN before installing software.",
        "Contact the helpdesk if setup fails.",
    ],
    must_not_contain=[],
)

print("=" * 60)
if FAILURES:
    print(f"{len(FAILURES)} test(s) FAILED: {FAILURES}")
    raise SystemExit(1)
else:
    print("All extraction tests passed.")
